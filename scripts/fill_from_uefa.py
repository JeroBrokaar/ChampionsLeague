"""Fill a results tab of the Europa/Conference League workbook from UEFA's match data.

Usage:
  python scripts/fill_from_uefa.py <workbook.xlsx> <tab> <competition> <seasonYear> [--write]
    competition: EL (Europa League) or ECL (Conference League)
    seasonYear:  UEFA's season year = the year the season ends (2025 for 2024/25)
  Without --write it only checks the matching and prints a report.

Scores: Home/Away goals = result after extra time (UEFA "total"); Penalties winner from the shoot-out.
Only empty cells are filled; scores already typed in are never changed.
"""
import json, os, re, sys, unicodedata, urllib.request
from collections import defaultdict
from datetime import datetime, timedelta

from openpyxl import load_workbook

COMPETITION_IDS = {"EL": 14, "ECL": 2019}
ROUNDS = {"League Phase", "Knockout Phase Play-Offs", "Knock-out Play-off", "Round of 16", "Quarter-finals", "Semi-finals", "Final"}


def uefa_matches(comp, season):
    url = f"https://match.uefa.com/v5/matches?competitionId={COMPETITION_IDS[comp]}&seasonYear={season}&limit=500&offset=0&order=ASC"
    with urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"}), timeout=60) as r:
        data = json.load(r)
    return [m for m in data if m.get("round", {}).get("metaData", {}).get("name") in ROUNDS]


def norm(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    s = re.sub(r"\b(fc|cf|ac|afc|sc|sk|fk|club|de|the|ssc|as|ss|kv|nk|bk|if|ff|cfr|1)\b", " ", s)
    return " ".join(s.split())


def similar(a, b):
    a, b = norm(a), norm(b)
    return a == b or (a and b and (a in b or b in a)) or (a.split()[:1] == b.split()[:1] and len(a.split()[0]) > 3)


def main(path, tab, comp, season, write):
    if write and os.path.exists(os.path.join(os.path.dirname(path), "~$" + os.path.basename(path))):
        sys.exit("The workbook is open in Excel - close it first, then run again.")
    wb = load_workbook(path)
    ws = wb[tab]
    rows = [(i, str(ws.cell(row=i, column=1).value), ws.cell(row=i, column=3).value, ws.cell(row=i, column=4).value)
            for i in range(2, ws.max_row + 1) if ws.cell(row=i, column=3).value]
    games = uefa_matches(comp, season)

    # UEFA dates are in UTC; Wikipedia's are local. Allow the same day or the day before/after.
    by_date = defaultdict(list)
    for g in games:
        d = datetime.fromisoformat(g["kickOffTime"]["dateTime"].replace("Z", "+00:00")).date()
        for off in (-1, 0, 1):
            by_date[str(d + timedelta(days=off))].append(g)

    name_map, used, pairs = {}, set(), {}
    # Pass 1: rows where both names look alike; pass 2: one name already known, the other inferred.
    for strict in (True, False):
        for i, date, home, away in rows:
            if i in pairs:
                continue
            cands = []
            for g in by_date[date]:
                if g["id"] in used:
                    continue
                gh, ga = g["homeTeam"]["internationalName"], g["awayTeam"]["internationalName"]
                hm = name_map.get(home) == gh or (name_map.get(home) is None and similar(home, gh))
                am = name_map.get(away) == ga or (name_map.get(away) is None and similar(away, ga))
                if (hm and am) or (not strict and (name_map.get(home) == gh or name_map.get(away) == ga)):
                    cands.append(g)
            if len(cands) == 1:
                g = cands[0]
                pairs[i] = g
                used.add(g["id"])
                name_map.setdefault(home, g["homeTeam"]["internationalName"])
                name_map.setdefault(away, g["awayTeam"]["internationalName"])

    unmatched = [(i, d, h, a) for i, d, h, a in rows if i not in pairs]
    conflicts = defaultdict(set)
    for i, _, h, a in rows:
        if i in pairs:
            conflicts[h].add(pairs[i]["homeTeam"]["internationalName"])
            conflicts[a].add(pairs[i]["awayTeam"]["internationalName"])
    conflicts = {k: v for k, v in conflicts.items() if len(v) > 1}

    print(f"{tab}: {len(rows)} rows, {len(games)} UEFA matches, {len(pairs)} matched, {len(unmatched)} unmatched")
    renamed = sorted((k, v) for k, v in name_map.items() if norm(k) != norm(v))
    if renamed:
        print("  name differences:", "; ".join(f"{k} = {v}" for k, v in renamed))
    if unmatched:
        print("  UNMATCHED:", unmatched[:10])
    if conflicts:
        print("  CONFLICTS:", conflicts)
    if unmatched or conflicts or len(pairs) != len(games):
        sys.exit("Not writing: the matching isn't complete and unambiguous.")

    filled = pens = skipped = 0
    for i, g in pairs.items():
        if ws.cell(row=i, column=5).value is not None or ws.cell(row=i, column=6).value is not None:
            skipped += 1
            continue
        if g.get("status") != "FINISHED":
            continue
        total = g["score"]["total"]
        ws.cell(row=i, column=5).value = total["home"]
        ws.cell(row=i, column=6).value = total["away"]
        filled += 1
        p = g["score"].get("penalty")
        if p and p["home"] != p["away"]:
            ws.cell(row=i, column=7).value = ws.cell(row=i, column=3 if p["home"] > p["away"] else 4).value
            pens += 1
    print(f"  {'would fill' if not write else 'filled'} {filled} results ({pens} with a penalty shoot-out), skipped {skipped} already filled")
    if write:
        wb.save(path)
        print("  saved", path)


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if a != "--write"]
    main(args[0], args[1], args[2], int(args[3]), "--write" in sys.argv)
