"""Create an Excel workbook to fill in Europa League and Conference League results.

Fixtures (date, round, home, away) and league-phase pots come from Wikipedia.
Usage:  python scripts/uefa_results_workbook.py 2024 2025
        -> europa-league/uefa-results-2024-25_2025-26.xlsx  (one tab per competition and season)
"""
import os, re, sys
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(__file__))
from el_fixtures_from_wikipedia import KO_ROUNDS, matches, pots, wikitext  # noqa: E402

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "europa-league")
COMPETITIONS = [
    # key, tab prefix, Wikipedia name, league-phase games per team, pots
    ("EL", "Europa League", "UEFA_Europa_League", 8, 4),
    ("ECL", "Conference League", "UEFA_Conference_League", 6, 6),
]
COLUMNS = ["Date", "Round", "Home", "Away", "Home goals", "Away goals", "Penalties winner"]
FONT = "Arial"
INPUT_FILL = PatternFill("solid", start_color="FFF2CC")
HEAD_FILL = PatternFill("solid", start_color="1F2937")
THIN = Side(style="thin", color="D0D5DD")


def fixtures(wiki_name, year):
    """League phase + knockouts + final, in date order, plus the pots."""
    title = f"{year}–{str(year + 1)[-2:]}_{wiki_name}"
    league_text = wikitext(f"{title}_league_phase")
    rows = []
    for date, _, sub, home, away in matches(league_text):
        md = re.search(r"Matchday (\d)", sub)
        rows.append([date, f"League Stage {md.group(1) if md else '?'}", home, away])
    seen = defaultdict(int)
    for date, top, _, home, away in matches(wikitext(f"{title}_knockout_phase")):
        if top in KO_ROUNDS:
            key = (top, frozenset((home, away)))
            seen[key] += 1
            rows.append([date, f"{KO_ROUNDS[top]} - {'1st' if seen[key] == 1 else '2nd'} leg", home, away])
    final_title = f"{year + 1}_{wiki_name}_final"
    for date, _, _, home, away in matches(wikitext(final_title))[:1]:
        rows.append([date, "Final", home, away])
    placeholder = re.compile(r"(Winner|Loser|Runner-up)\b")
    rows = [r for r in rows if not placeholder.match(r[2]) and not placeholder.match(r[3])]
    return sorted(rows, key=lambda r: r[0]), pots(league_text)


def check(label, rows, pot_map, games, n_pots):
    """Sanity checks; returns a list of problems (empty = fine)."""
    problems = []
    league = [r for r in rows if r[1].startswith("League")]
    teams = {r[2] for r in league} | {r[3] for r in league}
    home, away = Counter(r[2] for r in league), Counter(r[3] for r in league)
    if len(teams) != 36:
        problems.append(f"{len(teams)} teams instead of 36")
    if len(league) != 36 * games // 2:
        problems.append(f"{len(league)} league-phase matches instead of {36 * games // 2}")
    bad = [t for t in teams if home[t] != games // 2 or away[t] != games // 2]
    if bad:
        problems.append(f"not {games // 2} home + {games // 2} away: {bad}")
    pot_of = {t: p for p, ts in pot_map.items() for t in ts}
    if len(pot_map) != n_pots or set(pot_of) != teams:
        problems.append(f"pots don't match the teams ({len(pot_map)} pots)")
    ko = Counter(r[1].split(" - ")[0] for r in rows if not r[1].startswith("League"))
    expected = {"Play-offs": 16, "Round of 16": 16, "Quarter-finals": 8, "Semi-finals": 4, "Final": 1}
    for rnd, n in expected.items():
        # Rounds not drawn yet (current season) are simply absent; a partial round is a problem.
        if ko.get(rnd, 0) not in (0, n):
            problems.append(f"{rnd}: {ko.get(rnd, 0)} matches instead of {n}")
    ko_teams = {t for r in rows if not r[1].startswith("League") for t in r[2:4]}
    if ko_teams - teams:
        problems.append(f"knockout names not in league phase: {sorted(ko_teams - teams)}")
    print(f"{label}: {len(league)} league-phase + {sum(ko.values())} knockout matches - " + ("OK" if not problems else "; ".join(problems)))
    return problems


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def results_sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    ws.append(COLUMNS)
    style_header(ws, len(COLUMNS))
    for r in rows:
        ws.append([r[0], r[1], r[2], r[3], None, None, None])
    last = len(rows) + 1
    goals = DataValidation(type="whole", operator="between", formula1="0", formula2="30", allow_blank=True,
                           error="Type a whole number of goals (0-30).", errorTitle="Goals")
    ws.add_data_validation(goals)
    goals.add(f"E2:F{last}")
    for i in range(2, last + 1):
        for c in range(1, 8):
            cell = ws.cell(row=i, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = Border(bottom=THIN)
            if c >= 5:
                cell.fill = INPUT_FILL
                cell.alignment = Alignment(horizontal="center")
        # Penalties: pick one of the two teams of this match
        pens = DataValidation(type="list", formula1=f"=$C${i}:$D${i}", allow_blank=True)
        ws.add_data_validation(pens)
        pens.add(f"G{i}")
    for col, width in zip("ABCDEFG", (12, 26, 24, 24, 12, 12, 22)):
        ws.column_dimensions[col].width = width
    ws.auto_filter.ref = f"A1:G{last}"
    return ws, last


def main(years):
    wb = Workbook()
    guide = wb.active
    guide.title = "How to fill in"
    sheets = []
    pot_rows = []
    for key, name, wiki, games, n_pots in COMPETITIONS:
        for y in years:
            label = f"{name} {y}-{str(y + 1)[-2:]}"
            rows, pot_map = fixtures(wiki, y)
            check(label, rows, pot_map, games, n_pots)
            tab = f"{key} {y}-{str(y + 1)[-2:]}"
            ws, last = results_sheet(wb, tab, rows)
            sheets.append((tab, label, last))
            for p in sorted(pot_map):
                pot_rows += [[label, int(p), t] for t in pot_map[p]]

    # Pots tab
    ws = wb.create_sheet("Pots")
    ws.append(["Competition", "Pot", "Team"])
    style_header(ws, 3)
    for r in pot_rows:
        ws.append(r)
        for c in range(1, 4):
            ws.cell(row=ws.max_row, column=c).font = Font(name=FONT, size=10)
    for col, width in zip("ABC", (30, 8, 26)):
        ws.column_dimensions[col].width = width

    # Guide tab: instructions, example row and progress per tab
    g = guide
    g.column_dimensions["A"].width = 30
    for col in "BCDEFG":
        g.column_dimensions[col].width = 16
    g["A1"] = "Europa League & Conference League results"
    g["A1"].font = Font(name=FONT, bold=True, size=14)
    lines = [
        "Fill in only the yellow cells: Home goals and Away goals (whole numbers).",
        "Use the score after extra time when a knockout match went to extra time.",
        "Penalties winner: only for a match decided on penalties - pick the team from the dropdown.",
        "Leave both goal cells empty for a match that hasn't been played.",
        "Don't change dates, rounds or team names. Fixtures and pots: Wikipedia.",
    ]
    for i, text in enumerate(lines, start=3):
        g.cell(row=i, column=1, value=f"• {text}").font = Font(name=FONT, size=10)
    g["A9"] = "Example (format only, not a real result)"
    g["A9"].font = Font(name=FONT, bold=True, size=10)
    for c, (h, v) in enumerate(zip(COLUMNS, ["2025-03-13", "Round of 16 - 2nd leg", "Home FC", "Away FC", 1, 1, "Away FC"]), start=1):
        head = g.cell(row=10, column=c, value=h)
        head.font = Font(name=FONT, bold=True, color="FFFFFF")
        head.fill = HEAD_FILL
        cell = g.cell(row=11, column=c, value=v)
        cell.font = Font(name=FONT, size=10, italic=True)
        if c >= 5:
            cell.fill = INPUT_FILL
    g["A13"] = "Progress"
    g["A13"].font = Font(name=FONT, bold=True, size=10)
    for c, h in enumerate(["Tab", "Matches", "Filled in", "To do"], start=1):
        cell = g.cell(row=14, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
    for i, (tab, label, last) in enumerate(sheets, start=15):
        g.cell(row=i, column=1, value=label)
        g.cell(row=i, column=2, value=f"=COUNTA('{tab}'!C2:C{last})")
        # A match counts as filled in when both goal cells have a number
        g.cell(row=i, column=3, value=f"=SUMPRODUCT(ISNUMBER('{tab}'!E2:E{last})*ISNUMBER('{tab}'!F2:F{last}))")
        g.cell(row=i, column=4, value=f"=B{i}-C{i}")
        for c in range(1, 5):
            g.cell(row=i, column=c).font = Font(name=FONT, size=10)
    total = 15 + len(sheets)
    g.cell(row=total, column=1, value="Total").font = Font(name=FONT, bold=True, size=10)
    for c, col in ((2, "B"), (3, "C"), (4, "D")):
        cell = g.cell(row=total, column=c, value=f"=SUM({col}15:{col}{total - 1})")
        cell.font = Font(name=FONT, bold=True, size=10)
    g["B10"].comment = Comment("Columns are the same on every results tab.", "Generator")

    label = "_".join(f"{y}-{str(y + 1)[-2:]}" for y in years)
    path = os.path.join(OUT_DIR, f"uefa-results-{label}.xlsx")
    os.makedirs(OUT_DIR, exist_ok=True)
    if os.path.exists(path):
        sys.exit(f"{path} already exists - not overwriting (it may contain your results).")
    wb.calculation.fullCalcOnLoad = True  # progress formulas are computed when Excel opens the file
    wb.save(path)
    print("Saved", path)


if __name__ == "__main__":
    main([int(a) for a in sys.argv[1:]] or [2024, 2025])
